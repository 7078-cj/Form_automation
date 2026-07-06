from datetime import datetime
from utils.generate import generate

from docx import Document
from utils.lib import (
    format_period,
)

from docx import Document


def find_table_by_headers(doc: Document, expected_headers):
    """
    Finds a table whose first row matches the expected headers.

    Args:
        doc: Document object.
        expected_headers: List of header names in order.

    Returns:
        Table object.

    Raises:
        ValueError if no matching table is found.
    """

    expected = [h.strip().lower() for h in expected_headers]

    for table in doc.tables:
        if not table.rows:
            continue

        headers = [
            cell.text.strip().lower()
            for cell in table.rows[0].cells
        ]

        if headers == expected:
            return table

    raise ValueError(
        f"Could not find a table with headers: {expected_headers}"
    )

def fill_participants_table(doc: Document, participants):
    table = find_table_by_headers(
        doc,
        [
            "Student Full Name",
            "Student ID Number",
            "Year & Section",
            "Emergency Contact Number"
        ]
    )

    # Add rows if there aren't enough
    while len(table.rows) - 1 < len(participants):
        table.add_row()

    # Fill the table
    for i, participant in enumerate(participants, start=1):
        table.cell(i, 0).text = participant.get("Student_Full_Name", "")
        table.cell(i, 1).text = participant.get("Student ID Number", "")
        table.cell(i, 2).text = participant.get("Year & Section", "")
        table.cell(i, 3).text = participant.get("Emergency Contact Number", "")

def generate_itinerary(details):
    prompt = f"""
You are preparing an academic educational visit itinerary.

Competition Information

Title:
{details['title']}

Description:
{details['description']}

Venue:
{details['venue']}

Duration:
{format_period(details['start_date'], details['end_date'])}

Return ONLY valid JSON.

Schema:

{{
    "itinerary":[
        {{
            "activity":"...",
            "competency":"...",
            "assessment":"..."
        }}
    ]
}}

Rules:

- Generate between 3 and 6 activities.
- Activities should follow the chronological flow of the event.
- The first activity should be registration/orientation if appropriate.
- The final activity should be awarding, reflection, or closing ceremony if applicable.
- Competencies should be measurable learning outcomes.
- Assessments should be realistic academic assessments.
- Return JSON only.
"""

    return generate(prompt)

from docx import Document


def fill_itinerary_table(doc: Document, itinerary):
    table = find_table_by_headers(
        doc,
        [
            "Activity Segment / Itinerary Stop",
            "Target Curriculum Competency / LO",
            "Method of Evaluation / Assessment"
        ]
    )

    while len(table.rows) - 1 < len(itinerary):
        table.add_row()

    for i, item in enumerate(itinerary, start=1):
        table.cell(i, 0).text = item["activity"]
        table.cell(i, 1).text = item["competency"]
        table.cell(i, 2).text = item["assessment"]

def generate_risk_assessment(details):
    """
    Generates a risk assessment table for the competition.
    """

    prompt = f"""
You are preparing a Risk Assessment Matrix for an academic activity.

Competition Information

Title:
{details['title']}

Description:
{details['description']}

Venue:
{details['venue']}

Duration:
{format_period(details['start_date'], details['end_date'])}

Return ONLY valid JSON.

Schema:

{{
    "risks":[
        {{
            "hazard":"...",
            "risk_level":"...",
            "preventative_protocol":"...",
            "contingency_strategy":"..."
        }}
    ]
}}

Rules:

- Generate between 3 and 6 risks.
- Risks should be realistic for this competition.
- Risk Level should be one of:
    - Low
    - Moderate
    - High
    - Low / Moderate
    - Moderate / High
- Preventative Protocol should be concise (1 sentence).
- Emergency Contingency Strategy should be concise (1 sentence).
- Return JSON only.
"""

    return generate(prompt)

def fill_risk_assessment_table(doc: Document, risks):
    table = find_table_by_headers(
        doc,
        [
            "Identified Hazard Area",
            "Risk Level",
            "Preventative Protocol",
            "Emergency Contingency Strategy"
        ]
    )

    while len(table.rows) - 1 < len(risks):
        table.add_row()

    for i, risk in enumerate(risks, start=1):
        table.cell(i, 0).text = risk.get("hazard", "")
        table.cell(i, 1).text = risk.get("risk_level", "")
        table.cell(i, 2).text = risk.get("preventative_protocol", "")
        table.cell(i, 3).text = risk.get("contingency_strategy", "")

def generate_executive_summary(details):
    """
    Generates an executive summary and operational justification
    for the compliance form.
    """

    prompt = f"""
You are preparing an Executive Summary and Operational Justification
for an academic off-campus activity request.

Competition Information

Title:
{details['title']}

Description:
{details['description']}

Venue:
{details['venue']}

Duration:
{format_period(details['start_date'], details['end_date'])}

Return ONLY valid JSON.

Schema:

{{
    "executive_summary":"..."
}}

Rules:

- Write 1 professional paragraph.
- Explain the purpose of the activity.
- Explain the educational value and expected learning outcomes.
- Explain why participation is important for BS Information System students.
- Mention the itinerary's role in achieving the objectives.
- Mention that proper supervision and risk mitigation measures will be implemented.
- Do NOT use bullet points.
- Return JSON only.
"""

    return generate(prompt)

